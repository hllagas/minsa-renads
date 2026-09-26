# -*- coding: utf-8 -*-
"""Genera el reporte Word de pruebas automatizadas de RENADS.

Uso:
    python scripts/generar_reporte_pruebas.py
Salida:
    docs/reporte_pruebas_automatizadas.docx
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

AZUL = RGBColor(0x1F, 0x4E, 0x79)
GRIS = RGBColor(0x59, 0x59, 0x59)
VERDE = RGBColor(0x2E, 0x7D, 0x32)

FECHA = "25 de septiembre de 2026"


def _titulo(doc: Document, texto: str) -> None:
    p = doc.add_heading(texto, level=1)
    for run in p.runs:
        run.font.color.rgb = AZUL


def _subtitulo(doc: Document, texto: str) -> None:
    p = doc.add_heading(texto, level=2)
    for run in p.runs:
        run.font.color.rgb = AZUL


def _tabla(doc: Document, encabezados: list[str], filas: list[list[str]]):
    tabla = doc.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Light Grid Accent 1"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tabla.rows[0].cells
    for i, texto in enumerate(encabezados):
        hdr[i].paragraphs[0].add_run(texto).bold = True
    for fila in filas:
        celdas = tabla.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = str(valor)
    return tabla


def main() -> None:
    doc = Document()

    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    # Portada
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("Reporte de Pruebas Automatizadas")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = AZUL

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "RENADS — Registro Nacional de Articulación Docencia-Servicio en Salud"
    )
    r.italic = True
    r.font.size = Pt(13)
    r.font.color.rgb = GRIS

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"MINSA (Perú) · API REST (Django 6 + DRF)\nFecha: {FECHA}")

    doc.add_paragraph()

    # 1. Resumen ejecutivo
    _titulo(doc, "1. Resumen ejecutivo")
    doc.add_paragraph(
        "Se implementó una suite de pruebas unitarias automatizadas sobre la API "
        "de RENADS, cubriendo tres de los cinco módulos del sistema. La estrategia "
        "usa el framework nativo de Django (TestCase / DRF APITestCase), datos de "
        "prueba (mock data) construidos vía ORM y unittest.mock para las dependencias "
        "externas (almacenamiento en Cloudflare R2, LibreOffice/PDF, correo SMTP, "
        "reloj del sistema). La cobertura se mide con coverage.py."
    )
    p = doc.add_paragraph()
    p.add_run("Total de pruebas ejecutadas: ").bold = True
    r = p.add_run("543 — todas en verde (0 fallos, 0 errores).")
    r.bold = True
    r.font.color.rgb = VERDE

    _tabla(
        doc,
        ["Módulo", "Pruebas", "Cobertura", "Estado"],
        [
            ["common", "179", "95%", "OK"],
            ["internados", "162", "93%", "OK"],
            ["convenios", "202", "88.5% (código fuente)", "OK"],
            ["TOTAL", "543", "≥ 80% en todos", "OK"],
        ],
    )
    doc.add_paragraph()
    doc.add_paragraph(
        "Los módulos actividades y calendario aún no se prueban porque están "
        "pendientes de implementación."
    )

    # 2. Alcance y metodología
    _titulo(doc, "2. Alcance y metodología")
    for item in [
        "Tipo de prueba: unitarias e integración de endpoints DRF. No hay pruebas de "
        "interfaz (el sistema es una API sin UI).",
        "Framework: django.test.TestCase y rest_framework.test.APITestCase (nativo, sin pytest).",
        "Datos de prueba: creados vía ORM en setUpTestData; sin factory-boy ni faker.",
        "Mocking: unittest.mock/patch para almacenamiento R2/boto3, generación de PDF "
        "(docxtpl + LibreOffice + pypdf), correo (backend locmem) y timezone.now.",
        "Cobertura mínima exigida: 80% por módulo. Reporte HTML generado con coverage.py "
        "(htmlcov/index.html).",
        "Cada objetivo cubre happy path, unhappy path y edge case.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # 3. Detalle por módulo
    _titulo(doc, "3. Resultados por módulo")

    _subtitulo(doc, "3.1 common (autenticación, perfiles, permisos, almacenamiento, auditoría)")
    doc.add_paragraph("179 pruebas · cobertura total 95%.")
    _tabla(
        doc,
        ["Archivo fuente", "Cobertura"],
        [
            ["models.py", "100%"],
            ["permissions.py", "100%"],
            ["selectors.py", "100%"],
            ["filters.py", "100%"],
            ["urls.py / urls_2fa.py", "100%"],
            ["services.py", "99%"],
            ["serializers.py", "96%"],
            ["views.py", "92%"],
            ["storage.py", "82%"],
            ["documentai.py", "71% (cliente real de Google Document AI, no ejercitado)"],
        ],
    )
    doc.add_paragraph(
        "Cubre: modelo UserSecurity/UserProfile, auditoría, documentos versionados, "
        "JWT y claims, OTP por correo, 2FA (activar/verificar/confirmar/deshabilitar), "
        "alta y edición de usuarios (RN-username), alcance institucional, permisos "
        "(IsSuperUser, IsInstitutionalMember, IsModuleEnabled, HasEntityScope) y "
        "endpoints de login, /me, cambio de contraseña y CRUD de usuarios/grupos.",
        style="List Bullet",
    )

    _subtitulo(doc, "3.2 internados (estudiantes, internados, tutores, rotaciones)")
    doc.add_paragraph("162 pruebas · cobertura total 93%.")
    _tabla(
        doc,
        ["Archivo fuente", "Cobertura"],
        [
            ["models.py", "100%"],
            ["filters.py", "100%"],
            ["permissions.py", "100%"],
            ["urls.py", "100%"],
            ["selectors.py", "96%"],
            ["views.py", "93%"],
            ["serializers.py", "92%"],
            ["services.py", "85%"],
        ],
    )
    doc.add_paragraph(
        "Cubre reglas de negocio: RN-16 (carga masiva en dos pasos), RN-19 (periodo "
        "académico vs. especialidad), RN-20 (registro por universidad con alcance), "
        "RN-21 (unicidad de interno vigente), RN-22 (onboarding del interno), RN-23 "
        "(estado de declaraciones juradas), RN-24 (universidades del tutor) y RN-25 "
        "(interno atado a la asignación por universidad). Reutiliza el test_carga_masiva.py "
        "preexistente sin modificarlo.",
        style="List Bullet",
    )

    _subtitulo(doc, "3.3 convenios (ciclo de vida, campos clínicos, adendas, catálogos, PDF)")
    doc.add_paragraph(
        "202 pruebas · cobertura de código fuente de la app 88.5% (todos los archivos ≥ 80%)."
    )
    _tabla(
        doc,
        ["Archivo fuente", "Cobertura"],
        [
            ["selectors.py", "100%"],
            ["filters.py", "100%"],
            ["urls.py", "100%"],
            ["models.py", "97%"],
            ["permissions.py", "95%"],
            ["mixins.py", "94%"],
            ["serializers.py", "89%"],
            ["pdf.py", "85% (LibreOffice/docxtpl/pypdf mockeados)"],
            ["services.py", "82%"],
            ["views.py", "80%"],
        ],
    )
    doc.add_paragraph(
        "Cubre: RN-1 (quién solicita Convenio Marco, DIRIS exenta), ciclo de vida y "
        "transiciones de estado (forward-only e idempotente), campos clínicos en sus dos "
        "sub-módulos (registro CONAPRES y asignación por Órgano Regional con regla de "
        "disponibilidad y acumulador), adendas de ampliación y vigencia efectiva, "
        "nomenclatura, gate de campos clínicos con resolución CONAPRES antes de "
        "suscripción, partes firmantes y composición por tipo/categoría, representantes "
        "e historial, carreras por facultad (RN-FC), catálogos CRUD y generación de PDF "
        "(con todas las dependencias de sistema mockeadas).",
        style="List Bullet",
    )

    # 4. Defectos detectados y corregidos
    _titulo(doc, "4. Defectos detectados y corregidos")
    doc.add_paragraph(
        "Las pruebas revelaron dos defectos reales de la aplicación, que fueron "
        "corregidos con confirmación del responsable (respetando el flujo SDD):"
    )

    _subtitulo(doc, "4.1 Grafo de migraciones roto en build limpio")
    doc.add_paragraph(
        "La migración common 0006 creaba una FK hacia convenios.organdirectory pero solo "
        "dependía de convenios 0051. En una base de datos de prueba creada desde cero, "
        "Django ordenaba el renombrado OrganDirectory → OrganicUnit (convenios 0054) antes "
        "de common 0006, y la FK no resolvía (ValueError). Bloqueaba TODA la suite de pruebas.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Corrección: se añadió run_before = [(\"convenios\", \"0052_university_entity_type\")] "
        "a common 0006. Sin cambio de esquema. Suite verde en build limpio.",
        style="List Bullet",
    )

    _subtitulo(doc, "4.2 Desajuste de tipo en el alcance por solicitante")
    doc.add_paragraph(
        "En ConventionScope.has_object_permission (apps/convenios/permissions.py), la rama de "
        "solicitante comparaba solicitante_id_objeto (int) contra las entidades del usuario "
        "(id normalizado a str). El desajuste hacía que un usuario que solo fuese solicitante "
        "nunca pasara el control de acceso a nivel de objeto, pese a que el selector sí listaba "
        "el convenio (inconsistencia entre lo listado y lo accesible).",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Corrección: se castea str(obj.solicitante_id_objeto) en la comparación, alineado con "
        "el selector y la regla de normalización a str del proyecto. Verificado con 202 pruebas "
        "en verde.",
        style="List Bullet",
    )

    # 5. Cómo reproducir
    _titulo(doc, "5. Cómo reproducir las pruebas")
    doc.add_paragraph("Desde la raíz del proyecto, en PowerShell:")
    codigo = (
        ".venv\\Scripts\\Activate.ps1\n\n"
        "# Ejecutar un módulo\n"
        "python manage.py test apps.common\n"
        "python manage.py test apps.internados\n"
        "python manage.py test apps.convenios\n\n"
        "# Cobertura + reporte HTML de un módulo\n"
        "coverage run --source=apps.convenios manage.py test apps.convenios\n"
        "coverage report -m\n"
        "coverage html   # genera htmlcov/index.html"
    )
    p = doc.add_paragraph()
    run = p.add_run(codigo)
    run.font.name = "Consolas"
    run.font.size = Pt(9)

    doc.add_paragraph()
    doc.add_paragraph(
        "El reporte HTML de cobertura queda disponible en htmlcov/index.html. La "
        "dependencia coverage.py está declarada en requirements-dev.txt "
        "(coverage==7.16.1). El agente automatizado de pruebas vive en "
        ".claude/agents/testing.md."
    )

    # 6. Pendientes
    _titulo(doc, "6. Pendientes")
    for item in [
        "Módulo actividades: pruebas pendientes hasta que se implemente el módulo.",
        "Módulo calendario: pruebas pendientes hasta que se implemente el módulo.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    salida = Path(__file__).resolve().parents[1] / "docs" / "reporte_pruebas_automatizadas.docx"
    salida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(salida))
    print(f"Documento generado: {salida}")


if __name__ == "__main__":
    main()
