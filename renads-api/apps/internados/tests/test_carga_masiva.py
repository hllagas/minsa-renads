"""Pruebas de la carga masiva de estudiantes (RN-16): pre-validación, resaltado de
celdas y generación de la trama con cuadros combinados dependientes."""

import io
from types import SimpleNamespace

import openpyxl
from django.test import TestCase

from apps.convenios.models import AcademicLevel, ProfessionalCareer, Specialty, Ubigeo
from apps.internados import services
from apps.internados.models import (
    IdentityDocumentType,
    InternshipPeriod,
    RelationshipType,
    Student,
)


def _obtener_desde_dict(datos: dict):
    """Construye la función `obtener(col)` que espera `validar_fila_estudiante`."""
    def obtener(col):
        valor = datos.get(col)
        return services._celda(valor)
    return obtener


class NombreRangoTests(TestCase):
    def test_nombre_rango_saneado(self):
        self.assertEqual(services._nombre_rango("SAN MARTIN"), "R_SAN_MARTIN")
        self.assertEqual(services._nombre_rango("LIMA", "LIMA"), "R_LIMA_LIMA")

    def test_formula_subst_equivale_al_nombre(self):
        # La fórmula Excel debe reproducir la MISMA sustitución que `_nombre_rango`.
        expr = services._formula_subst("$K2")
        self.assertIn("UPPER($K2)", expr)
        for ch in services._RANGO_SUBS:
            self.assertIn(f'"{ch}","_"', expr)


class GeneracionTramaTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        IdentityDocumentType.objects.get_or_create(codigo="DNI", defaults={"nombre": "DNI"})
        IdentityDocumentType.objects.get_or_create(codigo="CE", defaults={"nombre": "Carné de extranjería"})
        RelationshipType.objects.get_or_create(codigo="PADRE", defaults={"nombre": "Padre"})
        RelationshipType.objects.get_or_create(codigo="MADRE", defaults={"nombre": "Madre"})
        Ubigeo.objects.get_or_create(codigo="150101", defaults=dict(departamento="LIMA", provincia="LIMA", distrito="LIMA"))
        Ubigeo.objects.get_or_create(codigo="150102", defaults=dict(departamento="LIMA", provincia="LIMA", distrito="ANCON"))
        Ubigeo.objects.get_or_create(codigo="220101", defaults=dict(departamento="SAN MARTIN", provincia="MOYOBAMBA", distrito="MOYOBAMBA"))
        nivel, _ = AcademicLevel.objects.get_or_create(codigo="PREGRADO", defaults={"nombre": "Pregrado"})
        ProfessionalCareer.objects.get_or_create(nombre="Medicina RN16 Trama", defaults={"nivel_academico": nivel})

    def test_estructura_trama(self):
        data = services.generar_trama_excel(es_pregrado=True)
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertIn("Estudiantes", wb.sheetnames)
        self.assertIn("_listas", wb.sheetnames)
        self.assertEqual(wb["_listas"].sheet_state, "hidden")

        # Rangos con nombre: Departamentos + por depto + por (depto,prov) + catálogos.
        nombres = set(wb.defined_names.keys())
        self.assertIn("Departamentos", nombres)
        self.assertIn("TipoDocumento", nombres)
        self.assertIn("Parentesco", nombres)
        self.assertIn("CarreraProfesional", nombres)  # carrera desde tabla (F7.2)
        self.assertIn(services._nombre_rango("LIMA"), nombres)
        self.assertIn(services._nombre_rango("LIMA", "LIMA"), nombres)
        self.assertIn(services._nombre_rango("SAN MARTIN"), nombres)

        ws = wb["Estudiantes"]
        encabezados = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertTrue(any((e or "").startswith("correo personal") for e in encabezados))
        self.assertTrue(any((e or "").startswith("teléfono móvil") for e in encabezados))
        # Debe haber validaciones de datos (cuadros combinados + custom).
        self.assertGreaterEqual(len(ws.data_validations.dataValidation), 8)


class ValidacionFilaTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.dni, _ = IdentityDocumentType.objects.get_or_create(codigo="DNI", defaults={"nombre": "DNI"})
        cls.ce, _ = IdentityDocumentType.objects.get_or_create(codigo="CE", defaults={"nombre": "Carné de extranjería"})
        cls.padre, _ = RelationshipType.objects.get_or_create(codigo="PADRE", defaults={"nombre": "Padre"})
        nivel, _ = AcademicLevel.objects.get_or_create(codigo="PREGRADO", defaults={"nombre": "Pregrado"})
        cls.carrera = ProfessionalCareer.objects.create(nombre="Medicina RN16", nivel_academico=nivel)
        cls.periodo = InternshipPeriod.objects.create(codigo="2025-I-RN16", nombre="2025-I")
        Ubigeo.objects.get_or_create(codigo="150101", defaults=dict(departamento="LIMA", provincia="LIMA", distrito="LIMA"))
        cls.uni = SimpleNamespace(id=1, codigo_inei="X")

    def _validar(self, datos, *, periodo=None):
        return services.validar_fila_estudiante(
            obtener=_obtener_desde_dict(datos), usuario=SimpleNamespace(),
            universidad=self.uni, periodo_internado=periodo, ct_uni=0, es_admin=True,
            es_pregrado_trama=True, dnis_vistos=set(),
        )

    def test_acumula_multiples_errores(self):
        datos = {
            "tipo_documento": "DNI", "numero_documento": "1234567",  # 7 dígitos → inválido
            "nombres": "JUAN", "apellido_paterno": "PEREZ",
            "correo": "correo-malo", "telefono": "abc",
            "departamento": "LIMA", "provincia": "LIMA", "distrito": "XXNODIST",
            "carrera_profesional": "Medicina RN16",
        }
        _, errores = self._validar(datos, periodo=self.periodo)
        columnas = {e["columna"] for e in errores}
        self.assertIn("numero_documento", columnas)
        self.assertIn("correo", columnas)
        self.assertIn("telefono", columnas)
        self.assertIn("distrito", columnas)

    def test_longitud_documento_por_tipo(self):
        base = {"nombres": "A", "apellido_paterno": "B", "carrera_profesional": "Medicina RN16"}
        # DNI exige 8 dígitos.
        _, err_dni = self._validar({**base, "tipo_documento": "DNI", "numero_documento": "12345678"}, periodo=self.periodo)
        self.assertNotIn("numero_documento", {e["columna"] for e in err_dni})
        _, err_dni_mal = self._validar({**base, "tipo_documento": "DNI", "numero_documento": "123456789"}, periodo=self.periodo)
        self.assertIn("numero_documento", {e["columna"] for e in err_dni_mal})
        # CE exige 9 dígitos.
        _, err_ce = self._validar({**base, "tipo_documento": "CE", "numero_documento": "123456789"}, periodo=self.periodo)
        self.assertNotIn("numero_documento", {e["columna"] for e in err_ce})

    def test_documento_texto_ceros_izquierda_y_solo_digitos(self):
        base = {"nombres": "A", "apellido_paterno": "B", "carrera_profesional": "Medicina RN16"}
        # Ceros a la izquierda: DNI "08123456" (8 caracteres) es válido y se conserva como texto.
        valores, err = self._validar({**base, "tipo_documento": "DNI", "numero_documento": "08123456"}, periodo=self.periodo)
        self.assertNotIn("numero_documento", {e["columna"] for e in err})
        self.assertEqual(valores["numero_documento"], "08123456")
        # Cualquier carácter no numérico se rechaza.
        _, err_alfa = self._validar({**base, "tipo_documento": "DNI", "numero_documento": "0812345X"}, periodo=self.periodo)
        self.assertIn("numero_documento", {e["columna"] for e in err_alfa})

    def test_nota_rango_y_decimales(self):
        base = {"tipo_documento": "DNI", "numero_documento": "12345678", "nombres": "A",
                "apellido_paterno": "B", "carrera_profesional": "Medicina RN16"}
        # Fuera de rango.
        _, e1 = self._validar({**base, "nota_promedio_ponderado": "21"}, periodo=self.periodo)
        self.assertIn("nota_promedio_ponderado", {e["columna"] for e in e1})
        # Más de 4 decimales.
        _, e2 = self._validar({**base, "nota_promedio_ponderado": "16.55555"}, periodo=self.periodo)
        self.assertIn("nota_promedio_ponderado", {e["columna"] for e in e2})
        # No numérica.
        _, e3 = self._validar({**base, "nota_promedio_ponderado": "abc"}, periodo=self.periodo)
        self.assertIn("nota_promedio_ponderado", {e["columna"] for e in e3})
        # Válida (4 decimales, en rango).
        v4, e4 = self._validar({**base, "nota_promedio_ponderado": "16.5000"}, periodo=self.periodo)
        self.assertNotIn("nota_promedio_ponderado", {e["columna"] for e in e4})

    def test_nombres_direccion_a_mayusculas(self):
        datos = {
            "tipo_documento": "DNI", "numero_documento": "12345678",
            "nombres": "juan carlos", "apellido_paterno": "pérez", "apellido_materno": "gómez",
            "direccion": "av. lima 123", "contacto_emergencia_nombre": "maría pérez",
            "carrera_profesional": "Medicina RN16",
        }
        valores, _ = self._validar(datos, periodo=self.periodo)
        self.assertEqual(valores["nombres"], "JUAN CARLOS")
        self.assertEqual(valores["apellido_paterno"], "PÉREZ")
        self.assertEqual(valores["apellido_materno"], "GÓMEZ")
        self.assertEqual(valores["direccion"], "AV. LIMA 123")
        self.assertEqual(valores["contacto_emergencia_nombre"], "MARÍA PÉREZ")

    def test_fila_valida_sin_errores(self):
        datos = {
            "tipo_documento": "DNI", "numero_documento": "12345678",
            "nombres": "JUAN", "apellido_paterno": "PEREZ",
            "correo": "juan@mail.com", "telefono": "999888777",
            "departamento": "LIMA", "provincia": "LIMA", "distrito": "LIMA",
            "carrera_profesional": "Medicina RN16", "contacto_emergencia_parentesco": "Padre",
        }
        valores, errores = self._validar(datos, periodo=self.periodo)
        self.assertEqual(errores, [])
        self.assertEqual(valores["numero_documento"], "12345678")
        self.assertEqual(valores["contacto_emergencia_parentesco"], self.padre)


class ValidarTramaTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        IdentityDocumentType.objects.get_or_create(codigo="DNI", defaults={"nombre": "DNI"})
        RelationshipType.objects.get_or_create(codigo="PADRE", defaults={"nombre": "Padre"})
        nivel, _ = AcademicLevel.objects.get_or_create(codigo="PREGRADO", defaults={"nombre": "Pregrado"})
        ProfessionalCareer.objects.create(nombre="Medicina RN16b", nivel_academico=nivel)
        cls.periodo = InternshipPeriod.objects.create(codigo="2025-I-RN16t", nombre="2025-I")
        Ubigeo.objects.get_or_create(codigo="150101", defaults=dict(departamento="LIMA", provincia="LIMA", distrito="LIMA"))
        cls.uni = SimpleNamespace(id=1, codigo_inei="X")

    def _trama(self, filas):
        headers = [
            "tipo_documento", "numero_documento", "apellido_paterno", "nombres",
            "correo personal", "teléfono móvil", "departamento", "provincia", "distrito",
            "carrera_profesional", "contacto_emergencia_parentesco",
        ]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for fila in filas:
            ws.append(fila)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def test_no_escribe_y_resalta(self):
        buf = self._trama([
            ["DNI", "1234567", "PEREZ", "JUAN", "malo", "x", "LIMA", "LIMA", "NOEXISTE", "Medicina", "Padre"],
        ])
        res = services.validar_trama_estudiantes(
            archivo=buf, usuario=SimpleNamespace(is_superuser=True), universidad=self.uni,
            periodo_internado=self.periodo,
        )
        self.assertEqual(Student.objects.count(), 0)
        self.assertIn(2, res["errores_por_fila"])

        anotado = services.anotar_trama(
            archivo=buf, errores_por_fila=res["errores_por_fila"], indice=res["indice"],
        )
        wb = openpyxl.load_workbook(io.BytesIO(anotado))
        ws = wb.active
        celdas_rojas = [
            c.coordinate for row in ws.iter_rows(min_row=2, max_row=2) for c in row
            if c.fill is not None and c.fill.fgColor is not None and c.fill.fgColor.rgb == "00FFC7CE"
        ]
        self.assertTrue(celdas_rojas, "Debe resaltar al menos una celda con inconsistencia")
